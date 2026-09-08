from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_BREAK


OUT = "documentacion/Como funciona el anonimizador.docx"
NAVY = "17324D"
BLUE = "DCEAF5"
PALE = "F4F7FA"
GRAY = "D9D9D9"
TEXT = RGBColor(36, 45, 54)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_cell_borders(cell, color=GRAY, size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        el = borders.find(qn(tag))
        if el is None:
            el = OxmlElement(tag)
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def keep_with_next(paragraph):
    paragraph.paragraph_format.keep_with_next = True


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].width = Inches(widths[i])
        hdr[i].text = h
        set_cell_shading(hdr[i], NAVY)
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(hdr[i])
        set_cell_borders(hdr[i])
        for run in hdr[i].paragraphs[0].runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.2)
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for r_idx, row in enumerate(rows):
        new_row = table.add_row()
        tr_pr = new_row._tr.get_or_add_trPr()
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)
        cells = new_row.cells
        for i, value in enumerate(row):
            cells[i].width = Inches(widths[i])
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cells[i])
            set_cell_borders(cells[i])
            if r_idx % 2:
                set_cell_shading(cells[i], PALE)
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(9.2)
                    run.font.color.rgb = TEXT
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    p.add_run(text)
    return p


def numbered(doc, number, title, body):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.22)
    r = p.add_run(f"{number}.  {title}  ")
    r.bold = True
    p.add_run(body)
    return p


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
sec.top_margin = Inches(0.72)
sec.bottom_margin = Inches(0.68)
sec.left_margin = Inches(0.78)
sec.right_margin = Inches(0.78)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.7)
normal.font.color.rgb = TEXT
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12

for name, size, before, after in (("Title", 28, 0, 14), ("Heading 1", 18, 14, 7), ("Heading 2", 13, 10, 4)):
    st = styles[name]
    st.font.name = "Aptos Display"
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = RGBColor(0, 0, 0)
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    st.paragraph_format.keep_with_next = True

for list_name in ("List Bullet", "List Bullet 2", "List Number"):
    styles[list_name].font.name = "Aptos"
    styles[list_name].font.size = Pt(10.5)
    styles[list_name].font.color.rgb = TEXT
    styles[list_name].paragraph_format.space_after = Pt(4)

if "Small label" not in styles:
    label = styles.add_style("Small label", WD_STYLE_TYPE.PARAGRAPH)
    label.font.name = "Aptos"
    label.font.size = Pt(9)
    label.font.bold = True
    label.font.color.rgb = RGBColor(60, 91, 118)
    label.font.all_caps = True
    label.paragraph_format.space_after = Pt(5)

# Cover
p = doc.add_paragraph(style="Small label")
p.add_run("GUÍA FUNCIONAL PARA EL EQUIPO")
p = doc.add_paragraph(style="Title")
p.add_run("Cómo funciona el anonimizador")
# Some Word renderers inherit a decorative border from the built-in Title style.
# Remove it explicitly so the title relies on typography and whitespace only.
p_pr = p._p.get_or_add_pPr()
p_bdr = p_pr.find(qn("w:pBdr"))
if p_bdr is not None:
    p_pr.remove(p_bdr)
p_style_pr = styles["Title"]._element.get_or_add_pPr()
p_style_bdr = p_style_pr.find(qn("w:pBdr"))
if p_style_bdr is not None:
    p_style_pr.remove(p_style_bdr)
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(18)
r = p.add_run("Explicación precisa del proyecto para personas sin formación técnica")
r.font.size = Pt(14)
r.font.color.rgb = RGBColor(67, 82, 96)

doc.add_paragraph(
    "El anonimizador prepara archivos con datos personales para poder trabajar con ellos sin exponer los valores reales. "
    "Busca información identificable, la sustituye por etiquetas coherentes y guarda en un archivo separado la relación necesaria para recuperarla. "
    "Todo el análisis se realiza en el equipo donde se ejecuta el programa."
)
p = doc.add_paragraph()
r = p.add_run("Idea principal  ")
r.bold = True
p.add_run(
    "El archivo anonimizado puede compartirse o procesarse con una herramienta de inteligencia artificial; el archivo de correspondencias debe permanecer bajo control, porque es la llave que permite reconstruir los datos originales."
)

doc.add_paragraph("Septiembre de 2026", style="Small label")
doc.add_page_break()

doc.add_heading("Resumen del funcionamiento", level=1)
doc.add_paragraph(
    "El proyecto cubre un recorrido completo de ida y vuelta. Primero crea una copia del archivo en la que nombres, correos, teléfonos, documentos de identidad y otros datos sensibles aparecen como etiquetas. Después se puede trabajar con esa copia. Por último, si las etiquetas se han conservado, el programa puede sustituirlas por sus valores originales."
)

numbered(doc, 1, "Anonimizar", "Lee el contenido, identifica datos personales y genera una copia con etiquetas como <PERSONA-1> o <EMAIL-1>.")
numbered(doc, 2, "Trabajar con la copia", "El archivo con etiquetas puede enviarse a una IA o compartirse con otra persona sin entregar directamente los valores sustituidos.")
numbered(doc, 3, "Restaurar", "Usa el archivo de correspondencias para volver a insertar cada valor original donde encuentre su etiqueta.")

doc.add_heading("Qué genera cada operación", level=2)
add_table(doc, ["Entrada", "Resultado principal", "Resultado auxiliar"], [
    ("informe.docx", "informe_anon.docx", "informe_anon.docx.key.json"),
    ("datos.csv", "datos_anon.csv", "datos_anon.csv.key.json"),
    ("resultado_ia.docx", "resultado_ia_restaurado.docx", "Usa un mapa ya existente"),
], [1.7, 2.35, 2.75])

doc.add_paragraph(
    "La etiqueta conserva la función del dato sin revelar su contenido. Si Juan García aparece cinco veces, todas las apariciones exactas reciben <PERSONA-1>. Esto permite mantener relaciones internas útiles: se entiende que se habla de la misma persona aunque no se conozca su nombre. La numeración vuelve a empezar para cada archivo procesado."
)

doc.add_heading("Qué datos intenta reconocer", level=1)
doc.add_paragraph(
    "Hay un grupo común que se busca con independencia del marco legal elegido y varios grupos específicos por país. El reconocimiento combina formas conocidas, como la estructura de un correo o de un documento nacional, con modelos de idioma capaces de reconocer nombres y lugares por el contexto."
)
add_table(doc, ["Grupo", "Ejemplos que busca", "Etiqueta visible"], [
    ("Común", "Personas, correos, teléfonos, IBAN, tarjetas, fechas, direcciones IP y lugares relacionados con personas", "<PERSONA-1>, <EMAIL-1>, <FECHA-1>"),
    ("España y Unión Europea", "DNI, NIE y teléfonos españoles", "<DNI-ES-1>"),
    ("Chile", "RUT o RUN y teléfonos chilenos", "<RUT-CL-1>"),
    ("Brasil", "CPF, CNPJ y teléfonos brasileños", "<CPF-BR-1>, <CNPJ-BR-1>"),
    ("México", "CURP, RFC y teléfonos mexicanos", "<CURP-MX-1>, <RFC-MX-1>"),
    ("Colombia", "NIT, cédula con una palabra indicativa y teléfonos colombianos", "<NIT-CO-1>, <DNI-ES-1>"),
    ("Argentina", "DNI, CUIT o CUIL y teléfonos argentinos", "<DNI-AR-1>, <CUIT-AR-1>"),
    ("Reino Unido", "National Insurance Number y teléfonos británicos", "<NINO-UK-1>"),
    ("California", "SSN, permiso de conducir y teléfonos estadounidenses", "<SSN-US-1>, <DL-US-1>"),
], [1.45, 3.55, 1.8])

doc.add_heading("Cómo decide qué sustituir", level=1)
doc.add_paragraph(
    "El programa no trata todos los textos del mismo modo. Primero determina el idioma del fragmento cuando hay suficiente texto. Después reúne las posibles coincidencias, elimina algunos errores conocidos y resuelve los casos en que dos detecciones ocupan las mismas letras. En un solapamiento conserva la opción con mayor grado de confianza. Los correos reciben prioridad para evitar que una parte del correo se interprete como otro tipo de dato."
)

doc.add_heading("Nombres, mayúsculas y textos breves", level=2)
doc.add_paragraph(
    "Para reconocer nombres usa los modelos de idioma instalados. Si más de la mitad de las palabras alfabéticas están en mayúsculas, crea una versión con mayúsculas normalizadas solo para el análisis, pero sustituye el texto original. También descarta saludos, tratamientos y palabras frecuentes que suelen confundirse con personas. En fragmentos de menos de cuatro palabras usa el primer idioma disponible, porque una muestra tan corta no permite identificar el idioma con fiabilidad."
)

doc.add_heading("La regla especial para lugares", level=2)
doc.add_paragraph(
    "Una ciudad aislada no siempre identifica a una persona. Para reducir sustituciones innecesarias, un lugar solo se oculta cuando existe una persona en el contexto relevante. En una línea de Markdown o un párrafo de Word, la persona debe quedar aproximadamente a 120 caracteres o menos. En una hoja de cálculo o CSV, basta con que haya una persona en cualquier celda de la misma fila. Así, la ciudad de una ficha de cliente se oculta, pero la región de una lista de productos puede conservarse."
)

doc.add_heading("Direcciones postales", level=2)
doc.add_paragraph(
    "Además de los lugares, existe una regla para direcciones escritas con tipos de vía habituales en español, como Calle, Avenida, Plaza, Paseo, Camino o Carretera, seguidos de un nombre y un número. Esta regla tiene prioridad sobre una detección parcial del lugar para intentar ocultar la dirección completa. Los formatos poco habituales o de otros idiomas pueden quedar fuera."
)

doc.add_heading("Cómo trata cada tipo de archivo", level=1)
add_table(doc, ["Formato", "Unidad de trabajo", "Consecuencia práctica"], [
    ("CSV", "Fila completa", "Analiza todas las celdas y usa la fila como contexto para los lugares."),
    ("Excel XLSX", "Cada fila de cada hoja", "Procesa todas las hojas. Reescribe el libro a partir de sus tablas de datos."),
    ("Markdown MD", "Cada línea", "Conserva el texto y los signos de Markdown, pero el contexto no cruza de una línea a otra."),
    ("Word DOCX", "Cada párrafo y las tablas", "Incluye cuerpo, encabezados, pies y cuadros de texto. No cubre notas al pie, notas finales ni comentarios."),
], [1.05, 2.05, 3.7])

doc.add_heading("Tratamiento particular de Word", level=2)
doc.add_paragraph(
    "Word puede dividir un nombre en varios fragmentos internos aunque visualmente parezca una sola frase. El anonimizador une el párrafo antes de analizarlo, por lo que puede detectar un nombre partido. Si no encuentra datos personales, deja intacta su composición. Si sustituye algo, concentra el texto resultante en el primer fragmento del párrafo y vacía los demás. El párrafo conserva el estilo del primer fragmento, pero pueden perderse cambios parciales de negrita, cursiva u otros estilos dentro de ese párrafo."
)
doc.add_paragraph(
    "La restauración de Word es más limitada: recorre el cuerpo y las tablas, fragmento por fragmento. No recorre encabezados, pies ni cuadros de texto. Además, si una etiqueta queda dividida en varios fragmentos, puede no reconocerla. Por eso conviene revisar especialmente esos elementos antes de considerar restaurado un documento Word."
)

doc.add_heading("Tratamiento particular de Screaming Frog", level=2)
doc.add_paragraph(
    "Un CSV se considera exportado por Screaming Frog si contiene las columnas Address, Content Type y Status Code. En ese caso no procesa toda la tabla. Solo revisa campos de texto libre como títulos, descripciones, encabezados, palabras clave y fragmentos de resultados. Las direcciones web y los datos técnicos del rastreo quedan fuera de ese recorrido."
)

doc.add_heading("El archivo de correspondencias", level=1)
doc.add_paragraph(
    "Cada archivo anonimizado tiene su propio mapa. El mapa relaciona cada etiqueta con el valor real e incluye la fecha, el nombre del archivo de origen y los marcos legales usados. Dos mapas pueden contener <PERSONA-1> con significados distintos; por eso no deben mezclarse."
)
bullet(doc, "Sin cifrado, el mapa guarda los datos originales en texto legible.")
bullet(doc, "Con cifrado, el contenido sensible y el nombre del archivo de origen quedan protegidos por una contraseña.")
bullet(doc, "Si se pierde el mapa, no existe información suficiente para recuperar los valores.")
bullet(doc, "Si se pierde la contraseña de un mapa cifrado, el proyecto no ofrece una vía de recuperación.")

h = doc.add_heading("Medidas de protección incorporadas", level=2)
h.paragraph_format.page_break_before = True
doc.add_paragraph(
    "Al guardar un mapa, el programa añade la regla *.key.json al archivo que indica qué elementos no deben incluirse en un repositorio de versiones. También avisa si reconoce en la ruta nombres habituales de carpetas sincronizadas, como OneDrive, Dropbox, Google Drive, iCloud o Box Sync. Estas medidas reducen errores accidentales, pero no sustituyen los permisos de acceso, las copias seguras ni el cifrado."
)
doc.add_paragraph(
    "La contraseña puede darse mediante una variable del entorno de trabajo, escribirse de forma oculta cuando se usa una terminal interactiva o incluirse directamente en la orden. Esta última opción es la menos segura porque puede quedar visible en el historial y en la lista de procesos."
)

doc.add_heading("Cómo funciona la restauración", level=1)
doc.add_paragraph(
    "Restaurar no vuelve a analizar nombres ni documentos. Busca etiquetas con la forma esperada y las reemplaza usando el mapa. Esto hace que el arranque sea rápido y evita que una segunda interpretación cambie el resultado."
)
numbered(doc, 1, "Mapa indicado expresamente", "Si se proporciona una ruta de mapa, usa esa para todas las entradas.")
numbered(doc, 2, "Coincidencia exacta", "Si no se indicó mapa, busca junto al archivo otro cuyo nombre sea exactamente el nombre del archivo más .key.json.")
numbered(doc, 3, "Único mapa de la carpeta", "Si el nombre no coincide, usa el mapa si solo hay uno en la carpeta. Esto permite restaurar un resultado al que una IA haya cambiado el nombre.")
numbered(doc, 4, "Ambigüedad", "Si hay varios mapas y ninguno coincide, se detiene para evitar usar correspondencias incorrectas.")

doc.add_heading("Comprobaciones después de restaurar", level=2)
doc.add_paragraph(
    "El programa informa de dos situaciones. La primera son etiquetas incluidas en el mapa que ya no aparecen en el archivo, por ejemplo porque una IA las eliminó o modificó. La segunda son cadenas con forma de etiqueta que aparecen en el documento pero no existen en ese mapa. Si no ocurre ninguna, comunica cuántas etiquetas encontró y que no quedaron residuos."
)

doc.add_heading("Opciones que cambian el comportamiento", level=1)
add_table(doc, ["Decisión", "Qué permite", "Precaución"], [
    ("Marco legal", "Elegir uno, varios o todos los grupos de reglas nacionales.", "Activar todos aumenta la cobertura y también los falsos positivos en números."),
    ("Archivo de salida", "Dar un nombre exacto cuando solo se procesa un archivo.", "No es válido para varias entradas."),
    ("Carpeta de salida", "Reunir allí todos los resultados.", "Se crea si todavía no existe."),
    ("Idiomas cargados", "Limitar los modelos a español, inglés u otros para acelerar el inicio.", "Un idioma no cargado se analiza con el primero disponible."),
    ("Cifrado del mapa", "Proteger las correspondencias con una contraseña.", "Hay que conservar la contraseña por separado."),
], [1.45, 3.05, 2.3])

doc.add_paragraph(
    "Hay una diferencia relevante entre las dos formas previstas de uso. Al ejecutar directamente el programa sin indicar marco legal, el valor actual es RGPD. La guía integrada para usarlo desde un asistente, en cambio, indica que si el usuario no menciona una jurisdicción se activen todas. Por tanto, el resultado por defecto depende de cómo se inicie el proceso."
)

doc.add_heading("Qué significa elegir un marco legal", level=2)
doc.add_paragraph(
    "La elección activa grupos de formas de identificación propias de cada país. No realiza una evaluación jurídica del archivo, no determina si el tratamiento cumple una ley y no certifica que el resultado sea anónimo en sentido legal. Es una selección de reglas de búsqueda. Los datos comunes se buscan siempre."
)

doc.add_heading("Límites y riesgos que el equipo debe conocer", level=1)
bullet(doc, "Puede dejar pasar datos personales si están escritos de una forma no contemplada o si el modelo de idioma no los reconoce.")
bullet(doc, "Puede ocultar códigos que no son personales, sobre todo al activar todos los países y sus reglas numéricas menos estrictas.")
bullet(doc, "No comprueba que DNI, RUT, CPF u otros números sean auténticos mediante su dígito de control; reconoce principalmente su forma.")
bullet(doc, "Las fechas se tratan como dato personal incluso cuando una fecha aislada podría no identificar a nadie.")
bullet(doc, "En una carpeta solo se revisan los archivos del primer nivel; no entra en subcarpetas.")
bullet(doc, "Una carpeta puede contener salidas anteriores y estas no se excluyen automáticamente al anonimizar; conviene usar una carpeta de destino limpia.")
bullet(doc, "El programa muestra errores de un archivo y continúa con los demás. Una ejecución puede terminar después de haber procesado solo parte del lote.")
bullet(doc, "Los archivos Excel se reconstruyen como tablas de datos. Elementos ajenos a esas tablas, como formatos complejos, fórmulas, gráficos o macros, no forman parte de la lógica de conservación.")
bullet(doc, "El contenido de Word en notas al pie, notas finales y comentarios no se anonimiza. La restauración tampoco cubre todos los lugares que sí cubre la anonimización.")
bullet(doc, "Si el texto original ya contiene algo como <EMAIL-1>, puede confundirse con una etiqueta durante la restauración. El programa avisa, pero no bloquea el proceso.")

doc.add_heading("Uso recomendado en el equipo", level=1)
numbered(doc, 1, "Conservar el original", "Trabajar siempre sobre una copia y mantener el archivo fuente en una ubicación controlada.")
numbered(doc, 2, "Elegir el alcance", "Usar el país adecuado cuando el archivo contiene muchos códigos; reservar todos los marcos para casos en que la cobertura adicional compense la revisión posterior.")
numbered(doc, 3, "Cifrar el mapa", "Hacerlo siempre que vaya a permanecer en una carpeta compartida o sincronizada.")
numbered(doc, 4, "Revisar la copia", "Buscar datos evidentes que hayan quedado visibles y comprobar sustituciones innecesarias antes de compartirla.")
numbered(doc, 5, "Mantener juntas las piezas correctas", "Guardar el resultado anonimizado y su mapa de forma que no haya dudas sobre qué mapa corresponde a cada archivo, pero con permisos distintos cuando sea posible.")
numbered(doc, 6, "Proteger las etiquetas", "Al pedir cambios a una IA, indicar que no debe modificar etiquetas como <PERSONA-1>.")
numbered(doc, 7, "Revisar la restauración", "Atender los avisos de etiquetas ausentes o sin mapa y comprobar manualmente encabezados, pies y cuadros de texto en Word.")

doc.add_heading("Ejemplo completo", level=1)
doc.add_paragraph("Texto original")
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.25)
p.add_run("Ana Ruiz, con DNI 12345678Z, vive en Sevilla y usa ana@ejemplo.com.").italic = True
doc.add_paragraph("Texto anonimizado")
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.25)
p.add_run("<PERSONA-1>, con DNI <DNI-ES-1>, vive en <UBICACION-1> y usa <EMAIL-1>.").italic = True
doc.add_paragraph("Correspondencias guardadas")
add_table(doc, ["Etiqueta", "Valor original"], [
    ("<PERSONA-1>", "Ana Ruiz"),
    ("<DNI-ES-1>", "12345678Z"),
    ("<UBICACION-1>", "Sevilla"),
    ("<EMAIL-1>", "ana@ejemplo.com"),
], [2.3, 4.5])
doc.add_paragraph(
    "Si la IA conserva las etiquetas, el programa puede reinsertar los cuatro valores. Si cambia <PERSONA-1> por PERSONA 1, ese nombre no se recuperará automáticamente y el informe de restauración advertirá de que la etiqueta del mapa no apareció."
)

doc.add_heading("Cobertura de las pruebas incluidas", level=1)
doc.add_paragraph(
    "El proyecto incluye una batería de pruebas que ejecuta el programa sobre ejemplos controlados. Comprueba identificadores de las ocho jurisdicciones, correos y direcciones; verifica que los valores desaparezcan de la copia, que el mapa los conserve, que un recorrido de ida y vuelta recupere CSV, Markdown y Word, y que un mapa cifrado no pueda abrirse sin contraseña."
)
doc.add_paragraph(
    "También verifica varios casos especialmente delicados: un nombre dividido internamente en Word, una ciudad asociada a una persona en la misma fila, una ciudad sin persona que debe mantenerse y un número de pedido que no debe confundirse con un identificador bajo RGPD. Las pruebas no cubren todas las variantes de formato, todos los idiomas ni la restauración de encabezados, pies o cuadros de texto."
)

doc.add_heading("Conclusión operativa", level=1)
doc.add_paragraph(
    "El anonimizador es una barrera local y reversible para reducir la exposición de datos personales antes de compartir o procesar archivos. Su utilidad depende de tres hábitos: revisar lo detectado, proteger el mapa y confirmar la restauración. No debe tratarse como una garantía absoluta ni como una certificación legal."
)
doc.add_paragraph(
    "Para el flujo habitual del equipo, la práctica más segura es trabajar con una copia, escoger el marco legal que corresponda, cifrar el mapa, comprobar que no queden datos visibles y exigir que cualquier herramienta externa conserve exactamente las etiquetas."
)

# Footer with page field
for section in doc.sections:
    footer = section.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Anonimizador  |  Guía funcional  |  ")
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(105, 115, 125)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    p._p.append(fld)

# Metadata and save
doc.core_properties.title = "Cómo funciona el anonimizador"
doc.core_properties.subject = "Guía funcional no técnica"
doc.core_properties.author = "Equipo del proyecto Anonimizador"
doc.core_properties.keywords = "anonimización, privacidad, guía funcional"

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
doc.save(OUT)
print(OUT)
